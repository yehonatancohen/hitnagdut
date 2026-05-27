#!/usr/bin/env python3
"""Generate Excel xlsx from objections data, output as base64."""

import sys
import json
import base64
import io
import math



try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("openpyxl not installed. Run: pip install openpyxl")
    sys.exit(1)


# ─── Colors ──────────────────────────────────────────────────────────────────
HEADER_BG  = "1F4E79"   # dark blue header row
HEADER_FG  = "FFFFFF"   # white text
CYAN_BG    = "CCFFFF"   # cyan — one per objection (PDF)
YELLOW_BG  = "FFF2CC"   # soft warm yellow for low confidence/missing warnings

# ─── Column definitions: (header_text, width) ────────────────────────────────
COLUMNS = [
    ("מס'",             6),
    ("מגיש ההתנגדות",   22),
    ("בשם",             18),
    ("כתובת",           15),
    ("גוש/חלקה",        12),
    ("סעיף",            10),
    ("מהות",            80),
    ("נושא",            15),
    ("נספח/מיקום התיקון", 20),
    ("נותן המענה",      22),
    ("מענה",            40),
]

# ─── Topic → advisor mapping (for reference sheet) ───────────────────────────
TOPIC_MAPPING = [
    ("תוספת יחידות דיור וזכויות בנייה",       "פרוגרמה, שמאי",                              "תוספת יחידות, תמ\"א 38, שינוי ייעוד, היטל השבחה, תמורות לבעלים"),
    ("תנועה, חניה וגישה",                      "יועץ תנועה",                                 "מקומות חניה, גישה לנכס, עומס תנועה, צמתים"),
    ("תשתיות — מים, ביוב, ניקוז",             "יועץ תשתיות",                                "ניקוז שטחים, קיבולת ביוב, לחץ מים, תעלות"),
    ("נוף ושטחים פתוחים",                      "אדריכל נוף",                                 "עצי נוי, שטחים ירוקים, גינות ציבוריות, מדרכות"),
    ("סביבה, רעש וזיהום",                      "יועץ סביבה",                                 "רעש, זיהום אוויר, קרינה, פסולת, ריחות"),
    ("חקלאות ועצים",                            "אגרונום",                                    "קרקע חקלאית, עצי פרי, ניקוז חקלאי, אדמת נחל"),
    ("אדריכלות ועיצוב — מבנה, גובה, קווי בנין", "אדריכל",                                   "גובה בנייה, מרחקים, חזית, מס' קומות, נסיגות"),
    ("שמאות, ירידת ערך ופיצויים",              "שמאי",                                       "ירידת ערך נכס, פיצויי הפקעה, תשלומי איזון"),
    ("מדיניות תכנונית כללית",                  "רשות מקומית, ועדת תכנון",                    "ייעוד קרקע, תוכניות מתאר, עקרונות תכנון"),
    ("התחדשות עירונית — פינוי-בינוי / תמ\"א", "הרשות להתחדשות עירונית, פרוגרמה, שמאי",    "פינוי-בינוי, תמ\"א 38, הסכמות דיירים, תמורות"),
]


def make_fill(hex_color: str) -> PatternFill:
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")


def make_border() -> Border:
    thin = Side(style="thin", color="CCCCCC")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def row_height_for(text: str) -> float:
    """Estimate row height in points based on text length in the מהות column."""
    chars_per_line = 55
    lines = max(1, math.ceil(len(text or '') / chars_per_line))
    return max(15, min(300, lines * 14))


def build_objections_sheet(wb: Workbook, objections: list) -> None:
    ws = wb.active
    ws.title = "נושאי ההתנגדות"
    ws.sheet_view.rightToLeft = True

    border = make_border()

    # ── Header row (row 1) ───────────────────────────────────────────────────
    hdr_font  = Font(name="Arial", size=11, bold=True, color=HEADER_FG)
    hdr_fill  = make_fill(HEADER_BG)
    hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)

    for col_idx, (title, _) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.alignment = hdr_align
        cell.border    = border

    ws.row_dimensions[1].height = 30

    # ── Data rows ────────────────────────────────────────────────────────────
    cyan_fill    = make_fill(CYAN_BG)
    no_fill      = PatternFill(fill_type=None)

    meta_font    = Font(name="Arial", size=10, bold=True)
    clause_font  = Font(name="Arial", size=10)

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)
    align_right  = Alignment(horizontal="right",  vertical="center", wrap_text=True, readingOrder=2)
    align_top_r  = Alignment(horizontal="right",  vertical="top",    wrap_text=True, readingOrder=2)
    align_top_c  = Alignment(horizontal="center", vertical="top",    wrap_text=True, readingOrder=2)

    excel_row = 2

    for obj_idx, obj in enumerate(objections, 1):
        meta    = obj.get("meta", {})

        megish      = meta.get("megish", "")
        beshem      = meta.get("beshem", "")
        ktovet      = meta.get("ktovet", "")
        gush_chelka = meta.get("gush_chelka", "")

        # ── Cyan header row ──────────────────────────────────────────────────
        cyan_values = [obj_idx, megish, beshem, ktovet, gush_chelka, "", "", "", "", "", ""]
        for col_idx, val in enumerate(cyan_values, 1):
            cell = ws.cell(row=excel_row, column=col_idx, value=val if val != "" else None)
            cell.fill   = cyan_fill
            cell.border = border
            if col_idx == 1:
                cell.font      = meta_font
                cell.alignment = align_center
            elif col_idx <= 5:
                cell.font      = meta_font
                cell.alignment = align_right
            else:
                cell.font      = clause_font
                cell.alignment = align_center

        ws.row_dimensions[excel_row].height = 18
        excel_row += 1

        # ── Sections & Clause rows ───────────────────────────────────────────
        sections = obj.get("sections", [])
        for sec in sections:
            sec_num     = str(sec.get("section_number", "") or "")
            sec_title   = str(sec.get("section_title", "") or "")
            sec_summary = str(sec.get("section_summary", "") or "")
            sec_annex   = str(sec.get("section_annex", "") or "")
            clauses     = sec.get("clauses", [])

            # Collect unique gorems for the section (preserving order)
            seen_gorems = set()
            unique_gorems = []
            for c in clauses:
                g = str(c.get("gorem", "") or "").strip()
                if g and g not in ("None", "none") and g not in seen_gorems:
                    seen_gorems.add(g)
                    unique_gorems.append(g)
            combined_gorem = "\n".join(unique_gorems)

            # Check confidence flags to highlight yellow
            missed_some_clauses = sec.get("missed_some_clauses", False)
            low_confidence      = sec.get("confidence", "") == "low"
            highlight_yellow    = missed_some_clauses or low_confidence
            sec_fill            = make_fill(YELLOW_BG) if highlight_yellow else no_fill

            start_row = excel_row

            # 1. Section Header Row
            row_values = [None, None, None, None, None, sec_num, sec_title, sec_summary, sec_annex, combined_gorem, None]
            for col_idx, val in enumerate(row_values, 1):
                cell = ws.cell(row=excel_row, column=col_idx, value=val)
                cell.fill   = sec_fill
                cell.border = border
                if col_idx == 6:        # seif
                    cell.font      = Font(name="Arial", size=10, bold=True)
                    cell.alignment = align_center
                elif col_idx == 7:      # mahut (title)
                    cell.font      = Font(name="Arial", size=10, bold=True)
                    cell.alignment = align_top_r
                elif col_idx in (8, 9): # nose, annex
                    cell.font      = Font(name="Arial", size=10, bold=True)
                    cell.alignment = align_center
                elif col_idx == 10:     # gorem (combined)
                    cell.font      = Font(name="Arial", size=10, bold=True)
                    cell.alignment = align_center
                else:
                    cell.font      = clause_font
                    cell.alignment = align_center

            ws.row_dimensions[excel_row].height = row_height_for(sec_title)
            excel_row += 1

            # 2. Section Clause Rows
            for clause in clauses:
                text  = str(clause.get("text", "") or "")

                row_values = [None, None, None, None, None, sec_num, text, sec_summary, sec_annex, None, None]
                for col_idx, val in enumerate(row_values, 1):
                    cell = ws.cell(row=excel_row, column=col_idx, value=val)
                    cell.fill   = sec_fill
                    cell.border = border
                    if col_idx == 6:        # seif
                        cell.font      = clause_font
                        cell.alignment = align_top_c
                    elif col_idx == 7:      # mahut
                        cell.font      = clause_font
                        cell.alignment = align_top_r
                    elif col_idx in (8, 9): # nose, annex
                        cell.font      = clause_font
                        cell.alignment = align_top_c
                    else:
                        cell.font      = clause_font
                        cell.alignment = align_center

                ws.row_dimensions[excel_row].height = row_height_for(text)
                excel_row += 1

            # 3. Vertically merge Section, Subject, Annex, and Gorem cells
            end_row = excel_row - 1
            if end_row > start_row:
                # Merge 'סעיף' column (6)
                ws.merge_cells(start_row=start_row, start_column=6, end_row=end_row, end_column=6)
                ws.cell(row=start_row, column=6).alignment = align_center
                ws.cell(row=start_row, column=6).font = Font(name="Arial", size=10, bold=True)

                # Merge 'נושא' column (8)
                ws.merge_cells(start_row=start_row, start_column=8, end_row=end_row, end_column=8)
                ws.cell(row=start_row, column=8).alignment = align_center
                ws.cell(row=start_row, column=8).font = Font(name="Arial", size=10, bold=True)

                # Merge 'נספח/מיקום התיקון' column (9)
                ws.merge_cells(start_row=start_row, start_column=9, end_row=end_row, end_column=9)
                ws.cell(row=start_row, column=9).alignment = align_center
                ws.cell(row=start_row, column=9).font = Font(name="Arial", size=10, bold=True)

                # Merge 'נותן המענה' column (10)
                ws.merge_cells(start_row=start_row, start_column=10, end_row=end_row, end_column=10)
                ws.cell(row=start_row, column=10).alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True, readingOrder=2
                )
                ws.cell(row=start_row, column=10).font = Font(name="Arial", size=10, bold=True)

    # ── Column widths ────────────────────────────────────────────────────────
    for col_idx, (_, width) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"


def build_mapping_sheet(wb: Workbook) -> None:
    """Add a reference sheet: topic category → who should respond."""
    ws = wb.create_sheet(title="מיפוי נושאים לגורמי מענה")
    ws.sheet_view.rightToLeft = True

    border = make_border()

    hdr_font  = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    hdr_fill  = make_fill("1F4E79")
    hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)

    headers = ["קטגוריית נושא", "גורמי המענה", "דוגמאות לנושאים"]
    widths  = [38, 38, 60]

    for col_idx, (hdr, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=col_idx, value=hdr)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.alignment = hdr_align
        cell.border    = border
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    ws.row_dimensions[1].height = 28

    row_align = Alignment(horizontal="right", vertical="center", wrap_text=True, readingOrder=2)
    body_font = Font(name="Arial", size=10)

    for row_idx, (topic, gorems, examples) in enumerate(TOPIC_MAPPING, 2):
        fill = make_fill("EBF2F8") if row_idx % 2 == 0 else PatternFill(fill_type=None)
        for col_idx, val in enumerate([topic, gorems, examples], 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font      = body_font
            cell.fill      = fill
            cell.alignment = row_align
            cell.border    = border
        ws.row_dimensions[row_idx].height = 20

    # Title note above table
    ws.insert_rows(1)
    title_cell = ws.cell(row=1, column=1,
        value="טבלת עזר: מיפוי נושאי התנגדות לגורמי מענה — ניתן להתאים לפי הפרויקט")
    title_cell.font      = Font(name="Arial", size=12, bold=True, color="1F4E79")
    title_cell.alignment = Alignment(horizontal="right", vertical="center", readingOrder=2)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)
    ws.row_dimensions[1].height = 24


def build_workbook(objections: list) -> bytes:
    wb = Workbook()
    build_objections_sheet(wb, objections)
    build_mapping_sheet(wb)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main():
    global sys
    sys.stdout = open(sys.stdout.fileno(), mode='w', encoding='utf-8', buffering=1)
    sys.stderr = open(sys.stderr.fileno(), mode='w', encoding='utf-8', buffering=1)
    sys.stdin  = open(sys.stdin.fileno(),  mode='r', encoding='utf-8', buffering=1)

    raw = sys.stdin.read()
    try:
        data = json.loads(raw)
    except json.JSONDecodeError as e:
        sys.stderr.write(f"Invalid JSON input: {e}\n")
        sys.exit(1)

    objections = data.get("objections", [])
    if not objections:
        sys.stderr.write("No objections provided\n")
        sys.exit(1)

    xlsx_bytes = build_workbook(objections)
    print(base64.b64encode(xlsx_bytes).decode('ascii'))


if __name__ == '__main__':
    main()
