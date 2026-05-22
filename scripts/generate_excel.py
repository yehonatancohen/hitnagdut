#!/usr/bin/env python3
"""Generate Excel xlsx from objections data, output as base64."""

import sys
import json
import base64
import io
import math



try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("openpyxl not installed. Run: pip install openpyxl")
    sys.exit(1)


# ─── Colors ──────────────────────────────────────────────────────────────────
HEADER_BG  = "1F4E79"   # dark blue header row
HEADER_FG  = "FFFFFF"   # white text
CYAN_BG    = "CCFFFF"   # cyan — one per objection (PDF)
YELLOW_BG  = "FFF2CC"   # soft warm yellow for low confidence/missing warnings

# ─── Column definitions: (header_text, width) ────────────────────────────────
COLUMNS = [
    ("מס'",             6),
    ("מגיש ההתנגדות",   22),
    ("בשם",             18),
    ("כתובת",           15),
    ("גוש/חלקה",        12),
    ("סעיף",            10),
    ("מהות",            80),
    ("נושא",            15),
    ("נספח/מיקום התיקון", 20),
    ("נותן המענה",      20),
    ("מענה",            40),
]


def make_fill(hex_color: str) -> PatternFill:
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")


def make_border() -> Border:
    thin = Side(style="thin", color="CCCCCC")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def row_height_for(text: str) -> float:
    """Estimate row height in points based on text length in the מהות column."""
    chars_per_line = 55
    lines = max(1, math.ceil(len(text or '') / chars_per_line))
    return max(15, min(300, lines * 14))


def build_workbook(objections: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "נושאי ההתנגדות"
    ws.sheet_view.rightToLeft = True

    border = make_border()

    # ── Header row (row 1) ───────────────────────────────────────────────────
    hdr_font  = Font(name="Arial", size=11, bold=True, color=HEADER_FG)
    hdr_fill  = make_fill(HEADER_BG)
    hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)

    for col_idx, (title, _) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.alignment = hdr_align
        cell.border    = border

    ws.row_dimensions[1].height = 30

    # ── Data rows ────────────────────────────────────────────────────────────
    cyan_fill    = make_fill(CYAN_BG)
    no_fill      = PatternFill(fill_type=None)

    meta_font    = Font(name="Arial", size=10, bold=True)
    clause_font  = Font(name="Arial", size=10)

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)
    align_right  = Alignment(horizontal="right",  vertical="center", wrap_text=True, readingOrder=2)
    align_top_r  = Alignment(horizontal="right",  vertical="top",    wrap_text=True, readingOrder=2)
    align_top_c  = Alignment(horizontal="center", vertical="top",    wrap_text=True, readingOrder=2)

    excel_row = 2

    for obj_idx, obj in enumerate(objections, 1):
        meta    = obj.get("meta", {})

        megish     = meta.get("megish", "")
        beshem     = meta.get("beshem", "")
        ktovet     = meta.get("ktovet", "")
        gush_chelka = meta.get("gush_chelka", "")

        # ── Cyan header row ──────────────────────────────────────────────────
        cyan_values = [obj_idx, megish, beshem, ktovet, gush_chelka, "", "", "", "", "", ""]
        for col_idx, val in enumerate(cyan_values, 1):
            cell = ws.cell(row=excel_row, column=col_idx, value=val if val != "" else None)
            cell.fill   = cyan_fill
            cell.border = border
            if col_idx == 1:
                cell.font      = meta_font
                cell.alignment = align_center
            elif col_idx <= 5:
                cell.font      = meta_font
                cell.alignment = align_right
            else:
                cell.font      = clause_font
                cell.alignment = align_center

        ws.row_dimensions[excel_row].height = 18
        excel_row += 1

        # ── Sections & Clause rows ───────────────────────────────────────────
        sections = obj.get("sections", [])
        for sec in sections:
            sec_num = str(sec.get("section_number", "") or "")
            sec_title = str(sec.get("section_title", "") or "")
            sec_summary = str(sec.get("section_summary", "") or "")
            sec_annex = str(sec.get("section_annex", "") or "")
            clauses = sec.get("clauses", [])

            # Check confidence flags to highlight yellow
            missed_some_clauses = sec.get("missed_some_clauses", False)
            low_confidence = sec.get("confidence", "") == "low"
            highlight_yellow = missed_some_clauses or low_confidence
            sec_fill = make_fill(YELLOW_BG) if highlight_yellow else no_fill

            start_row = excel_row

            # 1. Section Header Row
            row_values = [None, None, None, None, None, sec_num, sec_title, sec_summary, sec_annex, "", None]
            for col_idx, val in enumerate(row_values, 1):
                cell = ws.cell(row=excel_row, column=col_idx, value=val)
                cell.fill   = sec_fill
                cell.border = border
                if col_idx == 6:          # seif
                    cell.font      = Font(name="Arial", size=10, bold=True)
                    cell.alignment = align_center
                elif col_idx == 7:        # mahut (title)
                    cell.font      = Font(name="Arial", size=10, bold=True)
                    cell.alignment = align_top_r
                elif col_idx in (8, 9):   # nose (summary), annex
                    cell.font      = Font(name="Arial", size=10, bold=True)
                    cell.alignment = align_center
                else:
                    cell.font      = clause_font
                    cell.alignment = align_center

            ws.row_dimensions[excel_row].height = row_height_for(sec_title)
            excel_row += 1

            # 2. Section Clause Rows
            for clause in clauses:
                text = str(clause.get("text", "") or "")
                gorem = str(clause.get("gorem", "") or "")

                if gorem in ("None", "none"): gorem = ""

                row_values = [None, None, None, None, None, sec_num, text, sec_summary, sec_annex, gorem, None]
                for col_idx, val in enumerate(row_values, 1):
                    cell = ws.cell(row=excel_row, column=col_idx, value=val)
                    cell.fill   = sec_fill
                    cell.border = border
                    if col_idx == 6:          # seif
                        cell.font      = clause_font
                        cell.alignment = align_top_c
                    elif col_idx == 7:        # mahut
                        cell.font      = clause_font
                        cell.alignment = align_top_r
                    elif col_idx in (8, 9, 10):   # nose, annex, gorem
                        cell.font      = clause_font
                        cell.alignment = align_top_c
                    else:
                        cell.font      = clause_font
                        cell.alignment = align_center

                ws.row_dimensions[excel_row].height = row_height_for(text)
                excel_row += 1

            # 3. Vertically merge Section, Subject, and Annex cells
            end_row = excel_row - 1
            if end_row > start_row:
                # Merge 'סעיף' column (6)
                ws.merge_cells(start_row=start_row, start_column=6, end_row=end_row, end_column=6)
                ws.cell(row=start_row, column=6).alignment = align_center
                ws.cell(row=start_row, column=6).font = Font(name="Arial", size=10, bold=True)

                # Merge 'נושא' column (8)
                ws.merge_cells(start_row=start_row, start_column=8, end_row=end_row, end_column=8)
                ws.cell(row=start_row, column=8).alignment = align_center
                ws.cell(row=start_row, column=8).font = Font(name="Arial", size=10, bold=True)

                # Merge 'נספח/מיקום התיקון' column (9)
                ws.merge_cells(start_row=start_row, start_column=9, end_row=end_row, end_column=9)
                ws.cell(row=start_row, column=9).alignment = align_center
                ws.cell(row=start_row, column=9).font = Font(name="Arial", size=10, bold=True)

    # ── Column widths ────────────────────────────────────────────────────────
    for col_idx, (_, width) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main():
    # Force UTF-8 on Windows
    global sys
    sys.stdout = open(sys.stdout.fileno(), mode='w', encoding='utf-8', buffering=1)
    sys.stderr = open(sys.stderr.fileno(), mode='w', encoding='utf-8', buffering=1)
    sys.stdin  = open(sys.stdin.fileno(),  mode='r', encoding='utf-8', buffering=1)

    raw = sys.stdin.read()
    try:
        data = json.loads(raw)
    except json.JSONDecodeError as e:
        sys.stderr.write(f"Invalid JSON input: {e}\n")
        sys.exit(1)

    objections = data.get("objections", [])
    if not objections:
        sys.stderr.write("No objections provided\n")
        sys.exit(1)

    xlsx_bytes = build_workbook(objections)
    print(base64.b64encode(xlsx_bytes).decode('ascii'))


if __name__ == '__main__':
    main()
